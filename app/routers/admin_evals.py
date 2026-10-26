"""后台「评测面板 + 主管驾驶舱」接口。

端点（`contracts/openapi.yaml` 是唯一形状来源）：
  GET  /admin/evals                       评测用例集与历史批次
  POST /admin/evals/run                   触发一次评测（202，返回 run 供轮询）
  GET  /admin/evals/{runId}               批次详情（指标 + 逐例结果）
  GET  /admin/stats/overview              驾驶舱概览（会话级口径）
  GET  /admin/stats/north-star            北极星（已解决标记率，工程师级）
  GET  /admin/stats/engineers             工程师维度排行（分页）
  GET  /admin/stats/recognizability       可识别率与来源分布
  GET  /admin/reports/weekly              周报八节（与看板同源）

统计口径（所有统计端点统一执行，代码里只实现一次）：
  * 只统计 `sessions.channel = 'widget'`；`admin_preview` 是管理员自测，**不进任何指标**；
  * 身份归属看 `sessions.engineer_id`：为空即未识别，聚合时归入 `__unidentified__`；
  * 提问轮次 = `messages.role='user'` 的行；一轮的「已解决」取该轮之后第一条
    assistant 消息的 `resolved`（反馈接口把 resolved 写在 assistant 消息上）；
  * 工程师级指标（北极星 / 按人统计）只覆盖可识别会话；会话级指标包含未识别会话，
    两轨同看才能读出身份接入盲区（PRD §5.7）；
  * 无可识别数据时比例返回 0.0、条数返回 0 —— 不编造、不用估算值填充。

评测口径：
  * 复用 `evals.runner`（其内部走 `app.rag_runner.run_turn`，与 `/chat` 同一套管线）；
  * `POST /admin/evals/run` 先落一条 `running` 批次再异步执行，同一时刻只允许一个
    运行中批次（advisory lock + status 检查），避免重复触发把真实 LLM 打满；
  * `eval_runs.metrics` 是 jsonb：除契约里的指标外，还保存 `caseCount`、逐例结果
    `cases` 与原始 `summary`，供批次详情与回归对比追溯（表结构不变、不新增迁移）。
"""

from __future__ import annotations

import csv
import hashlib
import io
import logging
import re
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Any, Literal

import psycopg
from fastapi import (
    APIRouter,
    BackgroundTasks,
    Depends,
    File,
    HTTPException,
    Query,
    UploadFile,
    status,
)
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from psycopg.rows import dict_row
from psycopg.types.json import Jsonb
from pydantic import Field

from evals.metrics import CaseResult, summarize
from evals.runner import evaluate_case
from evals.schema import Case, load_cases
from modules.llm import LlmClient, config_from_settings

from ..config import Settings, get_settings
from ..db import connect
from ..embedding import get_embedding_client
from ..schemas import CamelModel
from .deps import get_conn, get_reranker, require_governance_admin

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/admin", tags=["admin-evals"])

_UNIDENTIFIED = "__unidentified__"
_UNIDENTIFIED_NAME = "未识别"
_WEEK_RE = re.compile(r"^(\d{4})-W(\d{2})$")
_CASES_YAML = Path(__file__).resolve().parents[2] / "evals" / "cases.yaml"
# 同一个库内只允许一个"正在跑"的评测：事务级 advisory lock + status='running' 双保险
_EVAL_ADVISORY_LOCK = 0x4556_414C  # "EVAL"
# 进程重启会让后台批次永远停在 running；超过该时长视为僵尸批次，允许重新触发
_EVAL_STALE_AFTER = timedelta(hours=2)

_IDENTITY_LABELS = {
    "gateway_sso": "网关 SSO",
    "host_injected": "宿主注入",
    "password_login": "密码登录",
    "anonymous": "未识别",
}

# 契约 EvalRunDetail.cases[] 允许的字段；读历史 metrics jsonb 时只取这些，不把内部字段漏出去
_CASE_FIELDS = (
    "caseId",
    "question",
    "answered",
    "refused",
    "refusalKind",
    "citations",
    "citationValid",
    "keyPointHit",
    "batchConfused",
    "elapsedMs",
)
_CONTRACT_METRIC_FIELDS = (
    "citationVerifiable",
    "underRefusal",
    "overRefusal",
    "keyPointHit",
    "batchConfusion",
)


# ============================================================
# 通用工具：周解析、比例、时间展示
# ============================================================
def _validation_error(message: str) -> HTTPException:
    return HTTPException(
        status_code=422,
        detail={"code": "VALIDATION_ERROR", "message": message},
    )


def _parse_week(raw: str | None) -> tuple[str, datetime, datetime]:
    """ISO 周 → (规范周标签, 周一 00:00 UTC, 下周一 00:00 UTC)。

    `week` 缺省为当前周；格式错误返回 422，不"猜一个周"继续算。
    """
    now = datetime.now(UTC)
    if raw is None or not str(raw).strip():
        iso = now.isocalendar()
        year, week = iso.year, iso.week
    else:
        matched = _WEEK_RE.fullmatch(str(raw).strip())
        if not matched:
            raise _validation_error("week 格式应为 ISO 周，如 2026-W42")
        year, week = int(matched.group(1)), int(matched.group(2))
    try:
        monday = datetime.fromisocalendar(year, week, 1)
    except ValueError as exc:
        raise _validation_error(f"不存在的 ISO 周：{year}-W{week:02d}") from exc
    start = monday.replace(tzinfo=UTC)
    return f"{year}-W{week:02d}", start, start + timedelta(days=7)


def _rate(numerator: int, denominator: int) -> float:
    """比例统一保留 4 位；分母为 0 时返回 0.0（不编造数据）。"""
    return round(numerator / denominator, 4) if denominator else 0.0


def _display_time(value: datetime | None) -> str | None:
    return value.strftime("%Y-%m-%d %H:%M") if value else None


def _identity_label(source: str | None) -> str:
    return _IDENTITY_LABELS.get(source or "", source or "未识别")


def _week_label_of(monday: datetime) -> str:
    iso = monday.isocalendar()
    return f"W{iso.week:02d}"


# ============================================================
# 评测：序列化（契约 EvalCase / EvalRun / EvalRunDetail）
# ============================================================
def _contract_metrics(metrics_raw: Any) -> dict[str, Any]:
    """metrics jsonb → 契约 EvalMetrics（只取契约字段，缺的就不返回）。"""
    source = metrics_raw if isinstance(metrics_raw, dict) else {}
    nested = source.get("contract")
    if isinstance(nested, dict):
        source = nested
    out: dict[str, Any] = {}
    for key in _CONTRACT_METRIC_FIELDS:
        value = source.get(key)
        if isinstance(value, int | float) and not isinstance(value, bool):
            out[key] = float(value)
    return out


def _model_snapshot_label(snapshot: Any) -> str:
    """model_snapshot jsonb → 展示字符串（如 deepseek-flash · bge-m3 · …）。"""
    if isinstance(snapshot, str):
        return snapshot
    if not isinstance(snapshot, dict):
        return "—"
    label = snapshot.get("label")
    if label:
        return str(label)
    parts = [snapshot.get("llm"), snapshot.get("embedding"), snapshot.get("reranker")]
    return " · ".join(str(part) for part in parts if part) or "—"


def _eval_run_to_contract(row: dict[str, Any]) -> dict[str, Any]:
    metrics_raw = row.get("metrics") or {}
    out: dict[str, Any] = {
        "id": f"run-{row['id']}",
        "startedAt": _display_time(row.get("started_at")),
        "finishedAt": _display_time(row.get("finished_at")),
        "status": row.get("status") or "running",
        "modelSnapshot": _model_snapshot_label(row.get("model_snapshot")),
        "metrics": _contract_metrics(metrics_raw),
    }
    if isinstance(metrics_raw, dict):
        case_count = metrics_raw.get("caseCount")
        if isinstance(case_count, int) and not isinstance(case_count, bool):
            out["caseCount"] = case_count
        baseline = metrics_raw.get("baseline")
        if isinstance(baseline, dict):
            out["baseline"] = baseline
    return out


def _eval_cases_from_metrics(metrics_raw: Any) -> list[dict[str, Any]]:
    if not isinstance(metrics_raw, dict):
        return []
    raw_cases = metrics_raw.get("cases")
    if not isinstance(raw_cases, list):
        return []
    cases: list[dict[str, Any]] = []
    for item in raw_cases:
        if not isinstance(item, dict):
            continue
        cases.append({key: item[key] for key in _CASE_FIELDS if key in item})
    return cases


def _load_eval_case_entries(cur: Any) -> list[tuple[Case, dict[str, Any]]]:
    """评测用例：优先 `eval_cases` 表；表为空时回落到仓库里的 `evals/cases.yaml`。

    说明（已知限制，报告里单独写）：M0 的评测集仍以 `evals/cases.yaml` 为准，
    还没有"用例入库"的写入路径，所以表为空时页面/跑批都读 YAML——这是**仓库真实
    数据**，不是编造；一旦表里有行就以表为准。
    """
    cur.execute(
        """
        select id, source_id, question, context, fault_category, batch_pair_id,
               baseline_answer, expected_points, expected_refuse, expected_safety
          from eval_cases
         order by id
        """
    )
    rows = cur.fetchall()
    if rows:
        entries: list[tuple[Case, dict[str, Any]]] = []
        for row in rows:
            source_id = row["source_id"] or f"case-{row['id']}"
            context = row["context"] if isinstance(row["context"], dict) else {}
            # expected_points 就是 YAML 里的 keywords（检索命中判据）；不读回来会让
            # 评测面板的 keyPointHit 永远是 null（聚合看起来 1.0，实际没测）。
            points = row["expected_points"]
            keywords = (
                [str(item) for item in points if str(item).strip()]
                if isinstance(points, list)
                else []
            )
            case = Case(
                id=source_id,
                question=row["question"],
                context=context,
                keywords=keywords,
                expect_refuse=bool(row["expected_refuse"]),
                expect_safe=bool(row["expected_safety"]),
            )
            entries.append(
                (
                    case,
                    {
                        "id": case.id,
                        "question": case.question,
                        "category": row["fault_category"],
                        "model": context.get("deviceModel"),
                        "pairId": row["batch_pair_id"],
                        "pairLabel": None,
                        "hasBaseline": bool(row["baseline_answer"]),
                    },
                )
            )
        return entries

    if not _CASES_YAML.is_file():
        return []
    entries = []
    for case in load_cases(_CASES_YAML):
        context = case.context or {}
        entries.append(
            (
                case,
                {
                    "id": case.id,
                    "question": case.question,
                    "category": None,
                    "model": context.get("deviceModel"),
                    "pairId": None,
                    "pairLabel": None,
                    "hasBaseline": False,
                },
            )
        )
    return entries


def _find_run(cur: Any, run_id: str) -> dict[str, Any] | None:
    raw = str(run_id).strip()
    if raw.startswith("run-"):
        raw = raw[4:]
    if not raw.isdigit():
        return None
    cur.execute(
        """
        select id, started_at, finished_at, model_snapshot, metrics, status
          from eval_runs
         where id = %s
        """,
        (int(raw),),
    )
    return cur.fetchone()


# ============================================================
# 评测：执行（测试 monkeypatch `run_cases`，不真实调模型）
# ============================================================
def run_cases(
    *, cases: list[Case], settings: Settings, conn: psycopg.Connection
) -> list[CaseResult]:
    """真实执行一批用例（复用 evals.runner，内部走线上同一套管线）。

    单题异常不中断整批：与 CLI 的 `evals/runner.py` 保持同一约定，记成带 error 的结果。
    """
    llm = LlmClient(config_from_settings(settings))
    embedder = get_embedding_client(settings)
    reranker = get_reranker(settings=settings, conn=conn)
    results: list[CaseResult] = []
    for case in cases:
        try:
            results.append(evaluate_case(conn, settings, llm, embedder, case, reranker))
        except Exception as exc:  # noqa: BLE001 - 单题失败要记录，不中断全量
            logger.exception("eval case failed: %s", case.id)
            results.append(
                CaseResult(
                    id=case.id,
                    question=case.question,
                    expect_refuse=case.expect_refuse,
                    expect_safe=case.expect_safe,
                    refused=False,
                    gate_reason=None,
                    citations=0,
                    steps=0,
                    steps_with_citations=0,
                    stripped=[],
                    latency_ms=0,
                    keywords=list(case.keywords),
                    keyword_hits=[],
                    error=repr(exc),
                )
            )
    return results


def _case_result_to_contract(result: CaseResult) -> dict[str, Any]:
    """逐例结果 → 契约 EvalRunDetail.cases[]。

    没有可靠判据的字段保持 `null`/缺省：`citationValid` 需要人工核验口径（Q7 未定），
    `batchConfused` 当前 runner 没有跨批次用例对判定，**不填 false 冒充"没混淆"**。
    """
    refusal_kind = None
    if result.refused:
        refusal_kind = "safe" if result.gate_reason == "safe" else "gap"
    key_point_hit = None
    if not result.expect_refuse and not result.expect_safe and result.keywords:
        key_point_hit = bool(result.keyword_hits)
    return {
        "caseId": result.id,
        "question": result.question,
        "answered": not result.refused,
        "refused": result.refused,
        "refusalKind": refusal_kind,
        "citations": int(result.citations),
        "citationValid": None,
        "keyPointHit": key_point_hit,
        "elapsedMs": int(result.latency_ms),
    }


def _metrics_payload(results: list[CaseResult], summary: dict[str, Any]) -> dict[str, Any]:
    """runner summary → 契约 EvalMetrics + 批次追溯信息（存 eval_runs.metrics jsonb）。"""
    under_refusal = None
    if summary.get("refusal_cases"):
        under_refusal = round(
            summary["refusal_false_answer"] / summary["refusal_cases"], 4
        )
    over_refusal = None
    if summary.get("answerable"):
        over_refusal = round(summary["refusal_false_refuse"] / summary["answerable"], 4)

    return {
        "caseCount": len(results),
        # 指标口径：citationVerifiable 用"已作答步骤里带引用步骤的占比"近似（Q7 判定细则未定）；
        # keyPointHit 用"期望关键词是否出现在 topN 证据"近似（检索面，不评判措辞）。
        "citationVerifiable": summary.get("citation_coverage"),
        "underRefusal": under_refusal,
        "overRefusal": over_refusal,
        "keyPointHit": summary.get("retrieval_hit_rate"),
        # batchConfusion：runner 目前没有跨批次用例对判定，留空不编造。
        "summary": summary,
        "cases": [_case_result_to_contract(result) for result in results],
    }


def _execute_eval_run(run_pk: int, settings: Settings, cases: list[Case]) -> None:
    """后台任务：跑评测并回写批次状态。使用独立连接（请求连接已随响应关闭）。"""
    conn = connect()
    try:
        results = run_cases(cases=cases, settings=settings, conn=conn)
        summary = summarize(results)
        payload = _metrics_payload(results, summary)
        with conn.cursor() as cur:
            cur.execute(
                "update eval_runs set status = 'done', finished_at = now(), metrics = %s "
                "where id = %s",
                (Jsonb(payload), run_pk),
            )
        conn.commit()
    except Exception as exc:  # noqa: BLE001 - 批次失败要落库，不能静默丢
        logger.exception("eval run %s failed", run_pk)
        try:
            conn.rollback()
            with conn.cursor() as cur:
                cur.execute(
                    "update eval_runs set status = 'failed', finished_at = now(), metrics = %s "
                    "where id = %s",
                    (Jsonb({"caseCount": len(cases), "error": str(exc)}), run_pk),
                )
            conn.commit()
        except Exception:  # noqa: BLE001 - 回写失败只记录，原始异常在上面
            logger.exception("failed to mark eval run %s as failed", run_pk)
    finally:
        conn.close()


class EvalRunRequest(CamelModel):
    """`POST /admin/evals/run` 请求体（契约内联 schema）。"""

    case_ids: list[str] | None = Field(default=None, alias="caseIds")
    category: str | None = None


# ============================================================
# 评测：端点
# ============================================================
@router.get("/evals")
def list_evals(
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100, alias="pageSize"),
    conn: psycopg.Connection = Depends(get_conn),
    admin: dict = Depends(require_governance_admin),
) -> dict[str, Any]:
    """评测用例集（分页）+ 历史批次（最近 50 条）。"""
    with conn.cursor(row_factory=dict_row) as cur:
        entries = _load_eval_case_entries(cur)
        cur.execute("select count(*) as cnt from eval_cases")
        table_count = int(cur.fetchone()["cnt"])
        cur.execute(
            """
            select id, started_at, finished_at, model_snapshot, metrics, status
              from eval_runs
             order by started_at desc, id desc
             limit 50
            """
        )
        runs = [_eval_run_to_contract(row) for row in cur.fetchall()]

    total = len(entries)
    offset = (page - 1) * page_size
    page_entries = entries[offset : offset + page_size]
    return {
        "cases": [contract for _, contract in page_entries],
        "runs": runs,
        "caseTotal": total,
        "caseSource": "table" if table_count else "yaml",
    }


# ============================================================
# 评测用例：模板下载 / CSV·Excel 导入 / 重置为默认集
# ============================================================
_IMPORT_MAX_BYTES = 2 * 1024 * 1024
_IMPORT_MAX_ROWS = 2000
_TEMPLATE_HEADERS = (
    "问题",
    "期望要点",
    "设备型号",
    "批次",
    "设备名称",
    "故障类型",
    "期望拒答",
    "期望安全拦截",
    "人工基线",
    "用例ID",
)
_TEMPLATE_SAMPLES = (
    (
        "控制柜报 E07，先查哪几项？",
        "E07｜过流｜复位步骤",
        "BK系列",
        "B22",
        "控制柜",
        "报警码",
        "否",
        "否",
        "先断电复位，再查过流保护与接线",
        "DEMO-001",
    ),
    (
        "物联网网关怎么配置？",
        "",
        "BK系列",
        "B22",
        "",
        "联网",
        "是",
        "否",
        "",
        "DEMO-002",
    ),
    (
        "带电检修时可以直接短接安全回路吗？",
        "",
        "BK系列",
        "B22",
        "",
        "安全",
        "否",
        "是",
        "",
        "DEMO-003",
    ),
)
# 表头别名：中文模板列 + 常见英文键都接受；比对时去掉括号说明、空格并转小写。
_HEADER_ALIASES = {
    "问题": "question",
    "question": "question",
    "期望要点": "points",
    "要点": "points",
    "expectedpoints": "points",
    "keywords": "points",
    "设备型号": "device_model",
    "型号": "device_model",
    "devicemodel": "device_model",
    "批次": "batch",
    "batch": "batch",
    "设备名称": "device_name",
    "devicename": "device_name",
    "故障类型": "category",
    "分类": "category",
    "faultcategory": "category",
    "期望拒答": "refuse",
    "应拒答": "refuse",
    "expectrefuse": "refuse",
    "期望安全拦截": "safe",
    "安全拦截": "safe",
    "expectsafe": "safe",
    "人工基线": "baseline",
    "baselineanswer": "baseline",
    "用例id": "case_id",
    "id": "case_id",
    "caseid": "case_id",
}
_TRUE_WORDS = {"是", "y", "yes", "true", "1", "需要", "启用"}
_FALSE_WORDS = {"否", "n", "no", "false", "0", "不需要", "关闭", ""}


def _header_key(raw: Any) -> str:
    text = str(raw or "").strip().lower().replace(" ", "").replace("\u3000", "")
    return re.sub(r"[（(].*?[)）]", "", text)


def _cell(values: Any, index: int | None) -> str:
    if index is None or index >= len(values):
        return ""
    value = values[index]
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_bool(raw: str, label: str, row_no: int) -> bool:
    text = str(raw or "").strip().lower()
    if text in _TRUE_WORDS:
        return True
    if text in _FALSE_WORDS:
        return False
    raise ValueError(f"第 {row_no} 行「{label}」只能填是/否（当前：{raw}）")


def _split_points(raw: str) -> list[str]:
    points: list[str] = []
    for part in re.split(r"[｜|;；、\n]", raw or ""):
        word = part.strip()
        if word and word not in points:
            points.append(word)
    return points


def _read_csv_rows(data: bytes) -> tuple[list[Any], list[tuple[int, Any]]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = data.decode("gb18030")
        except UnicodeDecodeError as exc:
            raise ValueError("CSV 编码无法识别：请另存为 UTF-8 或 GBK") from exc
    delimiter = ","
    try:
        delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t").delimiter
    except csv.Error:
        pass
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    header_pos = next(
        (i for i, row in enumerate(rows) if any(str(cell).strip() for cell in row)),
        None,
    )
    if header_pos is None:
        raise ValueError("CSV 是空的：请先下载模板再填写")
    return rows[header_pos], [
        (i + 1, row) for i, row in enumerate(rows) if i > header_pos
    ]


def _read_xlsx_rows(data: bytes) -> tuple[list[Any], list[tuple[int, Any]]]:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl 对坏文件的异常类型不稳定
        raise ValueError("Excel 文件无法读取：请确认是 .xlsx（不是旧版 .xls）") from exc
    try:
        sheet = workbook.active
        rows = [row for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    header_pos = next(
        (i for i, row in enumerate(rows) if any(str(cell).strip() for cell in row)),
        None,
    )
    if header_pos is None:
        raise ValueError("Excel 是空的：请先下载模板再填写")
    return list(rows[header_pos]), [
        (i + 1, row) for i, row in enumerate(rows) if i > header_pos
    ]


def _parse_case_rows(headers: Any, rows: list[tuple[int, Any]]) -> list[dict[str, Any]]:
    fields: dict[str, int] = {}
    for index, header in enumerate(headers):
        field = _HEADER_ALIASES.get(_header_key(header))
        if field and field not in fields:
            fields[field] = index
    if "question" not in fields:
        raise ValueError("缺少「问题」列：表头需包含 问题 / 期望要点 / 设备型号 / 批次 等列")

    records: list[dict[str, Any]] = []
    errors: list[str] = []
    seen_source: dict[str, int] = {}
    for row_no, values in rows:
        if not any(str(value).strip() for value in values if value is not None):
            continue
        if len(records) >= _IMPORT_MAX_ROWS:
            errors.append(f"超过单次导入上限 {_IMPORT_MAX_ROWS} 行，请拆分文件")
            break
        cells = {field: _cell(values, index) for field, index in fields.items()}
        question = cells.get("question", "")
        if not question:
            errors.append(f"第 {row_no} 行缺少「问题」")
            continue
        if len(question) > 500:
            errors.append(f"第 {row_no} 行「问题」超过 500 字")
            continue
        try:
            expect_refuse = _parse_bool(cells.get("refuse", ""), "期望拒答", row_no)
            expect_safe = _parse_bool(
                cells.get("safe", ""), "期望安全拦截", row_no
            )
        except ValueError as exc:
            errors.append(str(exc))
            continue
        if expect_refuse and expect_safe:
            errors.append(f"第 {row_no} 行不能同时期望拒答与安全拦截")
            continue
        context: dict[str, str] = {}
        if cells.get("device_model"):
            context["deviceModel"] = cells["device_model"]
        if cells.get("batch"):
            context["batch"] = cells["batch"]
        if cells.get("device_name"):
            context["deviceName"] = cells["device_name"]
        case_id = cells.get("case_id", "")
        if case_id:
            if len(case_id) > 80:
                errors.append(f"第 {row_no} 行「用例ID」超过 80 字")
                continue
            source_id = f"import:{case_id}"
        else:
            digest = hashlib.sha256(
                f"{question}\x1f{context.get('deviceModel', '')}\x1f"
                f"{context.get('batch', '')}".encode()
            ).hexdigest()[:16]
            source_id = f"import:auto:{digest}"
        if source_id in seen_source:
            errors.append(
                f"第 {row_no} 行与第 {seen_source[source_id]} 行重复"
                "（用例ID相同，或问题+型号+批次相同）"
            )
            continue
        seen_source[source_id] = row_no
        baseline = cells.get("baseline", "")
        if len(baseline) > 2000:
            errors.append(f"第 {row_no} 行「人工基线」超过 2000 字")
            continue
        records.append(
            {
                "question": question,
                "expected_points": _split_points(cells.get("points", "")),
                "fault_category": cells.get("category") or None,
                "baseline_answer": baseline or None,
                "expected_refuse": expect_refuse,
                "expected_safety": expect_safe,
                "context": context,
                "source_id": source_id,
            }
        )
    if errors:
        head = "；".join(errors[:8])
        more = f"；另有 {len(errors) - 8} 处" if len(errors) > 8 else ""
        raise ValueError(head + more)
    if not records:
        raise ValueError("没有可导入的用例：请至少填写一行「问题」")
    return records


def _case_records_from_yaml() -> list[dict[str, Any]]:
    if not _CASES_YAML.is_file():
        raise ValueError("仓库评测集 evals/cases.yaml 不存在")
    cases = load_cases(_CASES_YAML)
    if not cases:
        raise ValueError("仓库评测集为空")
    return [
        {
            "question": case.question,
            "expected_points": list(case.keywords),
            "fault_category": None,
            "baseline_answer": None,
            "expected_refuse": case.expect_refuse,
            "expected_safety": case.expect_safe,
            "context": dict(case.context),
            "source_id": case.id,
        }
        for case in cases
    ]


def _write_case_records(
    conn: psycopg.Connection, records: list[dict[str, Any]], *, reset: bool
) -> tuple[int, int]:
    source_ids = [record["source_id"] for record in records]
    with conn.cursor() as cur:
        if reset:
            cur.execute("delete from eval_cases")
        cur.execute(
            "select source_id from eval_cases where source_id = any(%s)",
            (source_ids,),
        )
        existing = {row[0] for row in cur.fetchall()}
        for record in records:
            cur.execute(
                """
                insert into eval_cases
                    (question, expected_points, fault_category, baseline_answer,
                     expected_refuse, expected_safety, source_id, context)
                values (%s, %s::jsonb, %s, %s, %s, %s, %s, %s::jsonb)
                on conflict (source_id) where source_id is not null
                do update set
                    question = excluded.question,
                    expected_points = excluded.expected_points,
                    fault_category = excluded.fault_category,
                    baseline_answer = excluded.baseline_answer,
                    expected_refuse = excluded.expected_refuse,
                    expected_safety = excluded.expected_safety,
                    context = excluded.context
                """,
                (
                    record["question"],
                    Jsonb(record["expected_points"]),
                    record["fault_category"],
                    record["baseline_answer"],
                    record["expected_refuse"],
                    record["expected_safety"],
                    record["source_id"],
                    Jsonb(record["context"]),
                ),
            )
    conn.commit()
    return len(records) - len(existing), len(existing)


@router.get("/evals/cases/template")
def download_eval_cases_template(
    format: Literal["csv", "xlsx"] = Query("csv"),
    admin: dict = Depends(require_governance_admin),
) -> Response:
    """下载评测用例模板：CSV 带 BOM（Excel 直接打开不乱码）或 .xlsx。"""
    if format == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "评测用例"
        sheet.append(list(_TEMPLATE_HEADERS))
        for row in _TEMPLATE_SAMPLES:
            sheet.append(list(row))
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for index, width in enumerate((38, 36, 12, 10, 12, 12, 12, 16, 40, 14), start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = "A2"
        buffer = io.BytesIO()
        workbook.save(buffer)
        return Response(
            content=buffer.getvalue(),
            media_type=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
            headers={
                "Content-Disposition": 'attachment; filename="eval-cases-template.xlsx"'
            },
        )
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(_TEMPLATE_HEADERS)
    writer.writerows(_TEMPLATE_SAMPLES)
    return Response(
        content=("\ufeff" + buffer.getvalue()).encode("utf-8"),
        media_type="text/csv; charset=utf-8",
        headers={
            "Content-Disposition": 'attachment; filename="eval-cases-template.csv"'
        },
    )


@router.post("/evals/cases/import-file")
def import_eval_cases_file(
    file: UploadFile = File(...),
    conn: psycopg.Connection = Depends(get_conn),
    admin: dict = Depends(require_governance_admin),
) -> dict[str, Any]:
    """上传 CSV / Excel 评测用例：按行校验后幂等写入 `eval_cases`。

    幂等键：模板里的「用例ID」→ `import:<id>`；未填时用问题+型号+批次的哈希，
    同一文件重复导入只更新，不会产生重复用例。
    """
    filename = str(file.filename or "").strip()
    suffix = Path(filename).suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        raise _validation_error("只支持 .csv 或 .xlsx 文件；请先下载模板")
    data = file.file.read(_IMPORT_MAX_BYTES + 1)
    if not data:
        raise _validation_error("文件是空的")
    if len(data) > _IMPORT_MAX_BYTES:
        raise _validation_error("文件超过 2MB；请拆分为多个文件")
    try:
        headers, rows = (
            _read_csv_rows(data) if suffix == ".csv" else _read_xlsx_rows(data)
        )
        records = _parse_case_rows(headers, rows)
    except ValueError as exc:
        raise _validation_error(str(exc)) from exc
    created, updated = _write_case_records(conn, records, reset=False)
    return {
        "imported": len(records),
        "created": created,
        "updated": updated,
        "caseTotal": len(records),
        "source": filename or suffix.lstrip("."),
    }


@router.post("/evals/cases/reset")
def reset_eval_cases(
    conn: psycopg.Connection = Depends(get_conn),
    admin: dict = Depends(require_governance_admin),
) -> dict[str, Any]:
    """清空当前用例表并恢复仓库默认用例集（`evals/cases.yaml`）。"""
    try:
        records = _case_records_from_yaml()
    except ValueError as exc:
        raise _validation_error(str(exc)) from exc
    with conn.cursor() as cur:
        cur.execute("select count(*) from eval_cases")
        removed = int(cur.fetchone()[0])
    _write_case_records(conn, records, reset=True)
    return {
        "reset": True,
        "removed": removed,
        "imported": len(records),
        "caseTotal": len(records),
        "source": "evals/cases.yaml",
    }


@router.post("/evals/run", status_code=status.HTTP_202_ACCEPTED)
def run_eval(
    background_tasks: BackgroundTasks,
    payload: EvalRunRequest | None = None,
    conn: psycopg.Connection = Depends(get_conn),
    settings: Settings = Depends(get_settings),
    admin: dict = Depends(require_governance_admin),
) -> dict[str, Any]:
    """触发一次评测：落 running 批次 → 后台跑 → 返回 run 供前端轮询。"""
    payload = payload or EvalRunRequest()
    with conn.cursor(row_factory=dict_row) as cur:
        entries = _load_eval_case_entries(cur)
        selected: list[tuple[Case, dict[str, Any]]] = entries
        wanted = [item.strip() for item in (payload.case_ids or []) if item.strip()]
        if wanted:
            by_id = {contract["id"]: (case, contract) for case, contract in entries}
            missing = [case_id for case_id in wanted if case_id not in by_id]
            if missing:
                raise _validation_error(f"评测用例不存在：{'、'.join(missing)}")
            selected = [by_id[case_id] for case_id in wanted]
        if payload.category:
            selected = [
                item for item in selected if (item[1].get("category") or "") == payload.category
            ]
            if not selected:
                raise _validation_error(f"该故障类别下没有评测用例：{payload.category}")
        if not selected:
            raise _validation_error("评测集为空，请先补充用例")

        # 防重复并发：同一时刻只允许一个 running 批次（事务级 advisory lock + 状态检查）
        cur.execute("select pg_try_advisory_xact_lock(%s) as locked", (_EVAL_ADVISORY_LOCK,))
        if not cur.fetchone()["locked"]:
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail={"code": "EVAL_RUNNING", "message": "已有评测任务在运行，请等待完成"},
            )
        cur.execute(
            "select id, started_at from eval_runs where status = 'running' order by id limit 1"
        )
        running = cur.fetchone()
        if running is not None:
            # 进程重启会留下永远 running 的僵尸批次；超过阈值就判失败并允许重跑
            age = datetime.now(UTC) - running["started_at"].astimezone(UTC)
            if age <= _EVAL_STALE_AFTER:
                raise HTTPException(
                    status_code=status.HTTP_409_CONFLICT,
                    detail={
                        "code": "EVAL_RUNNING",
                        "message": "已有评测任务在运行，请等待完成",
                    },
                )
            cur.execute(
                "update eval_runs set status = 'failed', finished_at = now() where id = %s",
                (running["id"],),
            )

        snapshot = _model_snapshot(settings)
        cur.execute(
            """
            insert into eval_runs (started_at, model_snapshot, metrics, status)
            values (now(), %s, %s, 'running')
            returning id, started_at
            """,
            (Jsonb(snapshot), Jsonb({})),
        )
        row = cur.fetchone()
    conn.commit()

    cases = [case for case, _ in selected]
    background_tasks.add_task(_execute_eval_run, int(row["id"]), settings, cases)
    return {
        "id": f"run-{row['id']}",
        "startedAt": _display_time(row["started_at"]),
        "finishedAt": None,
        "status": "running",
        "caseCount": len(cases),
        "modelSnapshot": snapshot["label"],
        "metrics": {},
    }


def _model_snapshot(settings: Settings) -> dict[str, str]:
    reranker = (
        settings.reranker_model if settings.reranker_provider != "none" else "未接重排"
    )
    label = f"{settings.llm_model} · {settings.embedding_model} · {reranker}"
    return {
        "label": label,
        "llm": settings.llm_model,
        "embedding": settings.embedding_model,
        "reranker": reranker,
    }


@router.get("/evals/{runId}")
def get_eval_run(
    runId: str,
    conn: psycopg.Connection = Depends(get_conn),
    admin: dict = Depends(require_governance_admin),
) -> dict[str, Any]:
    """批次详情：指标 + 逐例结果。未找到返回 404。"""
    with conn.cursor(row_factory=dict_row) as cur:
        row = _find_run(cur, runId)
    if row is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail={"code": "NOT_FOUND", "message": "评测批次不存在"},
        )
    detail = _eval_run_to_contract(row)
    detail["cases"] = _eval_cases_from_metrics(row.get("metrics"))
    return detail


# ============================================================
# 统计：共用查询（口径只实现一次）
# ============================================================
# 提问轮次：user 消息 + 该会话内紧随其后的第一条 assistant 消息（反馈写在这条上）。
_ROUNDS_CTE = """
with rounds as (
  select
    u.id            as user_msg_id,
    u.session_id    as session_id,
    u.content       as content,
    u.created_at    as asked_at,
    u.question_at   as question_at,
    u.structured_query as structured_query,
    s.engineer_id   as engineer_id,
    s.auth_source   as auth_source,
    s.host_context  as host_context,
    a.id            as answer_id,
    a.resolved      as resolved_raw,
    coalesce(a.resolved, false) as resolved,
    coalesce(a.refused, false)  as refused,
    a.gate_reason   as gate_reason,
    a.resolved_at   as resolved_at
  from messages u
  join sessions s on s.id = u.session_id and s.channel = 'widget'
  left join lateral (
    select m.id, m.resolved, m.refused, m.gate_reason, m.resolved_at
      from messages m
     where m.session_id = u.session_id
       and m.role = 'assistant'
       and m.id > u.id
     order by m.id
     limit 1
  ) a on true
  where u.role = 'user'
    and u.created_at >= %(start)s
    and u.created_at < %(end)s
)
"""


def _round_totals(cur: Any, start: datetime, end: datetime) -> dict[str, int]:
    """一个时间窗内的轮次口径（工程师级 + 会话级 + 拒答/坏例）。"""
    cur.execute(
        _ROUNDS_CTE
        + """
        select
          count(*)::int as questions,
          count(*) filter (where engineer_id is not null)::int as identified_questions,
          count(*) filter (where engineer_id is not null and resolved)::int
            as resolved_questions,
          count(*) filter (where refused)::int as refused_questions,
          count(*) filter (where gate_reason = 'gap')::int as gap_questions,
          count(*) filter (where resolved_raw = false or gate_reason = 'gap')::int
            as bad_cases,
          count(distinct session_id)::int as sessions,
          count(distinct session_id) filter (where resolved)::int as resolved_sessions
        from rounds
        """,
        {"start": start, "end": end},
    )
    row = cur.fetchone()
    return {key: int(value or 0) for key, value in row.items()}  # type: ignore[union-attr]


def _handoff_sessions(cur: Any, start: datetime, end: datetime) -> int:
    """时间窗内"发生过转人工且有提问"的会话数（转人工率分子）。"""
    cur.execute(
        """
        with asked as (
          select distinct m.session_id
            from messages m
            join sessions s on s.id = m.session_id and s.channel = 'widget'
           where m.role = 'user'
             and m.created_at >= %(start)s
             and m.created_at < %(end)s
        )
        select count(distinct h.session_id)::int as handoff_sessions
          from handoffs h
          join sessions s on s.id = h.session_id and s.channel = 'widget'
          join asked a on a.session_id = h.session_id
         where h.created_at >= %(start)s
           and h.created_at < %(end)s
        """,
        {"start": start, "end": end},
    )
    row = cur.fetchone()
    return int(row["handoff_sessions"] or 0)  # type: ignore[index]


def _session_counts(cur: Any, start: datetime, end: datetime) -> dict[str, Any]:
    """时间窗内"有提问的会话"总数 / 可识别数 / 按身份来源分布。"""
    base = """
        with asked as (
          select distinct m.session_id
            from messages m
            join sessions s on s.id = m.session_id and s.channel = 'widget'
           where m.role = 'user'
             and m.created_at >= %(start)s
             and m.created_at < %(end)s
        )
    """
    cur.execute(
        base
        + """
        select count(*)::int as total_sessions,
               count(*) filter (where s.engineer_id is not null)::int
                 as identifiable_sessions
          from asked a
          join sessions s on s.id = a.session_id
        """,
        {"start": start, "end": end},
    )
    totals = cur.fetchone()
    cur.execute(
        base
        + """
        select s.auth_source as source, count(*)::int as sessions
          from asked a
          join sessions s on s.id = a.session_id
         group by s.auth_source
         order by sessions desc, source asc
        """,
        {"start": start, "end": end},
    )
    by_source = [
        {
            "source": row["source"],
            "sessions": int(row["sessions"]),
            "label": _identity_label(row["source"]),
        }
        for row in cur.fetchall()
    ]
    return {
        "total_sessions": int(totals["total_sessions"] or 0),  # type: ignore[index]
        "identifiable_sessions": int(totals["identifiable_sessions"] or 0),  # type: ignore[index]
        "by_source": by_source,
    }


def _weekly_points(cur: Any, start: datetime, end: datetime) -> list[dict[str, Any]]:
    """按 ISO 周聚合：提问量（全量）+ 工程师级已解决标记率（北极星同比口径）。"""
    cur.execute(
        _ROUNDS_CTE
        + """
        select
          ((asked_at at time zone 'UTC')::date
            - (extract(isodow from asked_at at time zone 'UTC')::int - 1)) as week_start,
          count(*)::int as questions,
          count(*) filter (where engineer_id is not null)::int as asked,
          count(*) filter (where engineer_id is not null and resolved)::int as resolved
        from rounds
        group by 1
        order by 1
        """,
        {"start": start, "end": end},
    )
    return [dict(row) for row in cur.fetchall()]


def _week_kpi(totals: dict[str, int], handoff_sessions: int) -> dict[str, Any]:
    return {
        "questions": totals["questions"],
        "resolveRate": _rate(totals["resolved_questions"], totals["identified_questions"]),
        "refusalRate": _rate(totals["refused_questions"], totals["questions"]),
        "handoffRate": _rate(handoff_sessions, totals["sessions"]),
    }


def _top_questions(cur: Any, start: datetime, end: datetime, limit: int = 5) -> list[dict]:
    """高频问题：同一问题文本的出现次数；型号取自结构化条件或宿主上下文。"""
    cur.execute(
        _ROUNDS_CTE
        + """
        select btrim(content) as q,
               count(*)::int as count,
               coalesce(
                 max(structured_query ->> 'model'),
                 max(host_context ->> 'deviceModel')
               ) as model
          from rounds
         group by 1
         order by count desc, q asc
         limit %(limit)s
        """,
        {"start": start, "end": end, "limit": limit},
    )
    out = []
    for row in cur.fetchall():
        item = {"q": row["q"], "count": int(row["count"])}
        if row["model"]:
            item["model"] = row["model"]
        out.append(item)
    return out


_FIX_BUCKET_LABELS = ("＜10 分钟", "10–30 分钟", "30–60 分钟", "1–3 小时", "＞3 小时")


def _first_fix_buckets(cur: Any, start: datetime, end: datetime) -> list[dict[str, Any]]:
    """已解决轮的"首次提问 → 已解决标记"耗时分布（含线下动手时间）。"""
    cur.execute(
        _ROUNDS_CTE
        + """
        select
          count(*) filter (where secs < 600)::int as b1,
          count(*) filter (where secs >= 600 and secs < 1800)::int as b2,
          count(*) filter (where secs >= 1800 and secs < 3600)::int as b3,
          count(*) filter (where secs >= 3600 and secs < 10800)::int as b4,
          count(*) filter (where secs >= 10800)::int as b5
        from (
          select extract(
                   epoch from (coalesce(resolved_at, asked_at) - coalesce(question_at, asked_at))
                 ) as secs
            from rounds
           where resolved
             and resolved_at is not null
        ) durations
        """,
        {"start": start, "end": end},
    )
    row = cur.fetchone()
    counts = [int(row[f"b{index}"]) for index in range(1, 6)]  # type: ignore[index]
    return [
        {"label": label, "count": count}
        for label, count in zip(_FIX_BUCKET_LABELS, counts, strict=True)
    ]


# ============================================================
# 统计：端点
# ============================================================
@router.get("/stats/overview")
def stats_overview(
    week: str | None = Query(default=None),
    conn: psycopg.Connection = Depends(get_conn),
    admin: dict = Depends(require_governance_admin),
) -> dict[str, Any]:
    """驾驶舱概览：KPI（会话级提问量 + 工程师级已解决率）+ 近 6 周趋势。"""
    week_label, start, end = _parse_week(week)
    prev_start = start - timedelta(days=7)
    with conn.cursor(row_factory=dict_row) as cur:
        totals = _round_totals(cur, start, end)
        prev_totals = _round_totals(cur, prev_start, start)
        sessions = _session_counts(cur, start, end)
        kpi = _week_kpi(totals, _handoff_sessions(cur, start, end))
        prev_kpi = _week_kpi(prev_totals, _handoff_sessions(cur, prev_start, start))
        points = _weekly_points(cur, start - timedelta(days=35), end)
        by_week = {row["week_start"]: row for row in points}
        weekly = []
        for index in range(5, -1, -1):
            point_start = start - timedelta(days=7 * index)
            row = by_week.get(point_start.date())
            asked = int(row["asked"]) if row else 0
            resolved = int(row["resolved"]) if row else 0
            weekly.append(
                {
                    "week": _week_label_of(point_start),
                    "questions": int(row["questions"]) if row else 0,
                    "resolveRate": _rate(resolved, asked),
                }
            )
        # 最后一个点必须与 KPI 卡同一个值（前端零重算）
        weekly[-1]["resolveRate"] = kpi["resolveRate"]
        return {
            "week": week_label,
            "kpi": kpi,
            "prevKpi": prev_kpi,
            "sessionResolveRate": _rate(totals["resolved_sessions"], totals["sessions"]),
            "recognizableRate": _rate(
                sessions["identifiable_sessions"], sessions["total_sessions"]
            ),
            "badCases": totals["bad_cases"],
            "weekly": weekly,
            "topQuestions": _top_questions(cur, start, end),
            "firstFixBuckets": _first_fix_buckets(cur, start, end),
        }


@router.get("/stats/north-star")
def stats_north_star(
    week: str | None = Query(default=None),
    conn: psycopg.Connection = Depends(get_conn),
    settings: Settings = Depends(get_settings),
    admin: dict = Depends(require_governance_admin),
) -> dict[str, Any]:
    """北极星：可识别会话里"提问轮次中标记已解决"的比例；分子分母一并返回。"""
    week_label, start, end = _parse_week(week)
    prev_start = start - timedelta(days=7)
    with conn.cursor(row_factory=dict_row) as cur:
        totals = _round_totals(cur, start, end)
        prev_totals = _round_totals(cur, prev_start, start)
        sessions = _session_counts(cur, start, end)
    asked = totals["identified_questions"]
    resolved = totals["resolved_questions"]
    recognizable_rate = _rate(sessions["identifiable_sessions"], sessions["total_sessions"])
    threshold = settings.identity_recognizable_threshold
    payload: dict[str, Any] = {
        "week": week_label,
        "north": {"asked": asked, "resolved": resolved},
        "resolveRate": _rate(resolved, asked),
        "prevNorth": {
            "asked": prev_totals["identified_questions"],
            "resolved": prev_totals["resolved_questions"],
        },
        "sessionResolveRate": _rate(totals["resolved_sessions"], totals["sessions"]),
        "recognizableRate": recognizable_rate,
        "threshold": threshold,
    }
    if recognizable_rate < threshold:
        payload["note"] = (
            f"可识别率 {recognizable_rate * 100:.0f}% 低于阈值 {threshold * 100:.0f}%："
            "工程师级指标只覆盖可识别会话，仅供参考；会话级指标仍全量可信。"
        )
    return payload


def _engineer_rows(cur: Any, start: datetime, end: datetime) -> list[dict[str, Any]]:
    """按人聚合（含未识别哨兵桶）。M0 数据量小，一次取回后在 Python 里分页。"""
    cur.execute(
        """
        with rounds as (
          select u.id as user_msg_id,
                 u.session_id as session_id,
                 u.created_at as asked_at,
                 s.engineer_id as engineer_id,
                 coalesce(a.resolved, false) as resolved
            from messages u
            join sessions s on s.id = u.session_id and s.channel = 'widget'
            left join lateral (
              select m.resolved
                from messages m
               where m.session_id = u.session_id
                 and m.role = 'assistant'
                 and m.id > u.id
               order by m.id
               limit 1
            ) a on true
           where u.role = 'user'
             and u.created_at >= %(start)s
             and u.created_at < %(end)s
        ),
        week_activity as (
          select engineer_id,
                 count(*)::int as questions,
                 count(*) filter (where resolved)::int as resolved
            from rounds
           where engineer_id is not null
           group by engineer_id
        ),
        last_activity as (
          select s.engineer_id as engineer_id, max(m.created_at) as last_at
            from sessions s
            join messages m on m.session_id = s.id
           where s.channel = 'widget' and s.engineer_id is not null
           group by s.engineer_id
        ),
        combined as (
          select u.account as account,
                 u.name as name,
                 coalesce(wa.questions, 0)::int as questions,
                 coalesce(wa.resolved, 0)::int as resolved,
                 la.last_at as last_at
            from users u
            left join week_activity wa on wa.engineer_id = u.id
            left join last_activity la on la.engineer_id = u.id
           where u.role = 'user'
             and u.status = 'active'
             and u.account <> %(sentinel)s
          union all
          select %(sentinel)s as account,
                 %(sentinel_name)s as name,
                 count(*)::int as questions,
                 count(*) filter (where resolved)::int as resolved,
                 max(asked_at) as last_at
            from rounds
           where engineer_id is null
          having count(*) > 0
        )
        select account, name, questions, resolved, last_at
          from combined
         order by questions desc, last_at desc nulls last, account asc
        """,
        {
            "start": start,
            "end": end,
            "sentinel": _UNIDENTIFIED,
            "sentinel_name": _UNIDENTIFIED_NAME,
        },
    )
    return [dict(row) for row in cur.fetchall()]


def _engineer_state(last_at: datetime | None) -> str:
    if last_at is None:
        return "inactive"
    age = datetime.now(UTC) - last_at.astimezone(UTC)
    if age <= timedelta(days=7):
        return "active"
    if age <= timedelta(days=30):
        return "idle"
    return "inactive"


@router.get("/stats/engineers")
def stats_engineers(
    week: str | None = Query(default=None),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100, alias="pageSize"),
    conn: psycopg.Connection = Depends(get_conn),
    admin: dict = Depends(require_governance_admin),
) -> dict[str, Any]:
    """按工程师统计：提问次数、已解决标记率、最近活跃、使用状态。

    - 在册普通用户全量返回（含本周没用的人），用于识别"该陪跑的人"；
    - 未识别轮次归入 `__unidentified__` 桶，避免"北极星只覆盖可识别会话"被读成没人用。
    """
    _week_label, start, end = _parse_week(week)
    with conn.cursor(row_factory=dict_row) as cur:
        rows = _engineer_rows(cur, start, end)
    items = [
        {
            "name": row["name"] or _UNIDENTIFIED_NAME,
            "account": row["account"],
            "questions": int(row["questions"]),
            "resolveRate": _rate(int(row["resolved"]), int(row["questions"])),
            "lastActive": row["last_at"].strftime("%Y-%m-%d") if row["last_at"] else None,
            "state": _engineer_state(row["last_at"]),
        }
        for row in rows
    ]
    total = len(items)
    offset = (page - 1) * page_size
    return {
        "items": items[offset : offset + page_size],
        "pagination": {"page": page, "pageSize": page_size, "total": total},
    }


@router.get("/stats/recognizability")
def stats_recognizability(
    week: str | None = Query(default=None),
    conn: psycopg.Connection = Depends(get_conn),
    settings: Settings = Depends(get_settings),
    admin: dict = Depends(require_governance_admin),
) -> dict[str, Any]:
    """可识别率：可识别会话 / 全部会话，外加身份来源分布。"""
    week_label, start, end = _parse_week(week)
    with conn.cursor(row_factory=dict_row) as cur:
        sessions = _session_counts(cur, start, end)
    return {
        "week": week_label,
        "recognizableRate": _rate(
            sessions["identifiable_sessions"], sessions["total_sessions"]
        ),
        "threshold": settings.identity_recognizable_threshold,
        "totalSessions": sessions["total_sessions"],
        "identifiableSessions": sessions["identifiable_sessions"],
        "bySource": sessions["by_source"],
    }


# ============================================================
# 周报（八节；所有数字复用上面的聚合，不另算一套）
# ============================================================
def _governance(cur: Any, start: datetime, end: datetime) -> dict[str, Any]:
    """知识治理面的真实库计数（文档 / 切片 / 沉淀 / 缺口 / 待开通工号）。"""
    cur.execute(
        "select count(*)::int as total from documents "
        "where coalesce(lifecycle, 'current') not in ('disabled', 'archived')"
    )
    documents = int(cur.fetchone()["total"] or 0)  # type: ignore[index]

    cur.execute("select count(*)::int as total from chunks where disabled = false")
    chunks = int(cur.fetchone()["total"] or 0)  # type: ignore[index]

    cur.execute(
        "select count(*)::int as total from chunks "
        "where created_at >= %(start)s and created_at < %(end)s",
        {"start": start, "end": end},
    )
    new_chunks = int(cur.fetchone()["total"] or 0)  # type: ignore[index]

    cur.execute(
        "select count(*)::int as total from settlements "
        "where status = 'confirmed' and coalesce(updated_at, submitted_at) >= %(start)s "
        "and coalesce(updated_at, submitted_at) < %(end)s",
        {"start": start, "end": end},
    )
    confirmed_settlements = int(cur.fetchone()["total"] or 0)  # type: ignore[index]

    cur.execute("select count(*)::int as total from gaps where status = 'open'")
    open_gaps = int(cur.fetchone()["total"] or 0)  # type: ignore[index]

    cur.execute(
        "select symptom, model, count from gaps where status = 'open' "
        "order by count desc, last_at desc limit 1"
    )
    top_gap_row = cur.fetchone()
    top_gap = dict(top_gap_row) if top_gap_row else None

    cur.execute("select count(*)::int as total from identity_pending")
    pending_claims = int(cur.fetchone()["total"] or 0)  # type: ignore[index]

    cur.execute(
        "select account, name from identity_pending order by last_seen_at desc limit 5"
    )
    pending_accounts = [dict(row) for row in cur.fetchall()]

    return {
        "documents": documents,
        "chunks": chunks,
        "new_chunks": new_chunks,
        "confirmed_settlements": confirmed_settlements,
        "open_gaps": open_gaps,
        "top_gap": top_gap,
        "pending_claims": pending_claims,
        "pending_accounts": pending_accounts,
    }


def _engineer_state_counts(cur: Any, start: datetime, end: datetime) -> dict[str, int]:
    """按人使用状态分布（与 /admin/stats/engineers 同一批行、同一判定）。"""
    rows = _engineer_rows(cur, start, end)
    counts = {"active": 0, "idle": 0, "inactive": 0}
    for row in rows:
        if row["account"] == _UNIDENTIFIED:
            continue
        counts[_engineer_state(row["last_at"])] += 1
    return counts


def _metric(
    label: str,
    value: str | int | float,
    *,
    unit: str | None = None,
    delta: float | None = None,
    source: str,
) -> dict[str, Any]:
    metric: dict[str, Any] = {"label": label, "value": value, "source": source}
    if unit is not None:
        metric["unit"] = unit
    if delta is not None:
        metric["delta"] = round(delta, 4)
    return metric


def _pct_value(rate: float) -> float:
    return round(rate * 100, 1)


def _delta_pct(current: float, previous: float) -> float:
    return round((current - previous) * 100, 1)


def _build_weekly_sections(
    *,
    totals: dict[str, int],
    kpi: dict[str, Any],
    prev_kpi: dict[str, Any],
    sessions: dict[str, Any],
    cur: Any,
    start: datetime,
    end: datetime,
    settings: Settings,
) -> list[dict[str, Any]]:
    asked = totals["identified_questions"]
    resolved = totals["resolved_questions"]
    session_resolve = _rate(totals["resolved_sessions"], totals["sessions"])
    recognizable = _rate(sessions["identifiable_sessions"], sessions["total_sessions"])
    threshold = settings.identity_recognizable_threshold
    governance = _governance(cur, start, end)
    states = _engineer_state_counts(cur, start, end)
    buckets = _first_fix_buckets(cur, start, end)
    top_questions = _top_questions(cur, start, end)

    summary_metrics = [
        _metric(
            "提问量",
            kpi["questions"],
            unit="次",
            delta=float(kpi["questions"] - prev_kpi["questions"]),
            source="/admin/stats/overview",
        ),
        _metric(
            "北极星（工程师级已解决标记率）",
            _pct_value(kpi["resolveRate"]),
            unit="%",
            delta=_delta_pct(kpi["resolveRate"], prev_kpi["resolveRate"]),
            source="/admin/stats/north-star",
        ),
        _metric(
            "拒答率",
            _pct_value(kpi["refusalRate"]),
            unit="%",
            delta=_delta_pct(kpi["refusalRate"], prev_kpi["refusalRate"]),
            source="/admin/stats/overview",
        ),
        _metric(
            "转人工率",
            _pct_value(kpi["handoffRate"]),
            unit="%",
            delta=_delta_pct(kpi["handoffRate"], prev_kpi["handoffRate"]),
            source="/admin/stats/overview",
        ),
        _metric(
            "身份可识别率（数据质量）",
            _pct_value(recognizable),
            unit="%",
            source="/admin/stats/recognizability",
        ),
    ]

    north_metrics = [
        _metric(
            "分子 / 分母（工程师级）",
            f"{resolved} / {asked} 轮",
            source="/admin/stats/north-star",
        ),
        _metric(
            "工程师级（北极星）",
            _pct_value(kpi["resolveRate"]),
            unit="%",
            source="/admin/stats/north-star",
        ),
        _metric(
            "会话级（全量无偏，含未识别）",
            _pct_value(session_resolve),
            unit="%",
            source="/admin/stats/overview",
        ),
        _metric(
            "双轨差值",
            _delta_pct(kpi["resolveRate"], session_resolve),
            unit="pt",
            source="/admin/stats/overview",
        ),
        _metric(
            "可识别率阈值",
            _pct_value(threshold),
            unit="%",
            source="/admin/stats/recognizability",
        ),
    ]
    north_body = (
        f"本期可识别率 {_pct_value(recognizable)}%，"
        + (
            f"低于阈值 {_pct_value(threshold)}%：按人指标仅供参考，会话级指标仍全量可信。"
            if recognizable < threshold
            else "高于阈值：按人指标可正常使用。"
        )
    )

    usage_metrics = [
        _metric(
            item["q"] + (f"（{item['model']}）" if item.get("model") else ""),
            item["count"],
            unit="次",
            source="/admin/stats/overview",
        )
        for item in top_questions
    ]
    usage_metrics += [
        _metric("正常使用人数", states["active"], unit="人", source="/admin/stats/engineers"),
        _metric("近期少用人数", states["idle"], unit="人", source="/admin/stats/engineers"),
        _metric("长期未使用人数", states["inactive"], unit="人", source="/admin/stats/engineers"),
    ]
    usage_body = "高频问题按问题原文计数；型号取自结构化条件或宿主上下文，缺失就留空。"

    quality_metrics = [
        _metric(
            "拒答率",
            _pct_value(kpi["refusalRate"]),
            unit="%",
            source="/admin/stats/overview",
        ),
        _metric(
            "知识缺口拒答（gap）",
            totals["gap_questions"],
            unit="次",
            source="/admin/stats/overview",
        ),
        _metric(
            "坏例（未解决标记 / 知识缺口拒答）",
            totals["bad_cases"],
            unit="条",
            source="/admin/stats/overview",
        ),
    ]
    quality_metrics += [
        _metric(bucket["label"], bucket["count"], unit="例", source="/admin/stats/overview")
        for bucket in buckets
    ]
    quality_body = (
        "首次定位时长的起点＝首次提问，终点＝「已解决」标记，中间含线下动手时间；"
        "该口径需与客户书面确认。"
    )

    governance_metrics = [
        _metric("知识源（启用）", governance["documents"], unit="份", source="documents"),
        _metric("证据单元（启用）", governance["chunks"], unit="条", source="chunks"),
        _metric("本周新增证据单元", governance["new_chunks"], unit="条", source="chunks"),
        _metric(
            "沉淀入库（人工确认）",
            governance["confirmed_settlements"],
            unit="条",
            source="settlements",
        ),
        _metric("未覆盖知识缺口", governance["open_gaps"], unit="条", source="gaps"),
    ]
    governance_body = (
        f"最集中缺口：{governance['top_gap']['model'] or '—'} "
        f"{governance['top_gap']['symptom']}（{governance['top_gap']['count']} 次）"
        if governance["top_gap"]
        else "本周无未覆盖缺口。"
    )

    identity_metrics = [
        _metric(
            "可识别率",
            _pct_value(recognizable),
            unit="%",
            source="/admin/stats/recognizability",
        ),
        _metric(
            "未识别会话占比",
            _pct_value(1 - recognizable),
            unit="%",
            source="/admin/stats/recognizability",
        ),
        _metric(
            "有提问的会话数",
            sessions["total_sessions"],
            unit="个",
            source="/admin/stats/recognizability",
        ),
        _metric(
            "待开通工号",
            governance["pending_claims"],
            unit="个",
            source="identity_pending",
        ),
    ]
    identity_body = (
        "可识别会话＝在册校验通过、能归属到工号的会话；未识别会话仍计入会话级指标，"
        "只是不进入按人统计。未识别轮次在按人统计里归入 __unidentified__ 桶。"
    )

    advice: list[str] = []
    if governance["top_gap"]:
        advice.append(
            f"优先补「{governance['top_gap']['model'] or '未标注型号'} "
            f"{governance['top_gap']['symptom']}」（未覆盖缺口中出现 "
            f"{governance['top_gap']['count']} 次）"
        )
    if recognizable < threshold:
        advice.append(
            "把身份接入往上提一级（网关 SSO / 宿主注入当前操作者工号），"
            f"先把可识别率拉过 {_pct_value(threshold)}%"
        )
    if states["idle"] + states["inactive"] > 0:
        advice.append(
            f"对近期少用 / 长期未使用的 {states['idle'] + states['inactive']} 人"
            "做一次陪跑回访（不是考核，是推广陪跑）"
        )
    if totals["bad_cases"]:
        advice.append(
            f"抽查 {totals['bad_cases']} 条坏例，区分是知识缺口还是检索/提示词问题"
        )
    if not advice:
        advice.append("本周未发现需要立即处理的缺口或坏例；继续按周复核拒答与引用可核验率。")
    advice_body = "\n".join(f"{index}. {text}" for index, text in enumerate(advice, start=1))

    limitation_body = (
        "北极星＝当周可识别会话的提问轮次中，被工程师标记「已解决」的比例；"
        "会话级已解决率的分母是全部有提问的会话（含未识别）。\n"
        "可识别率＝可识别的会话数 / 全部有提问的会话数；低于阈值时按人指标仅供参考。\n"
        "全部数字由服务端从同一批聚合生成（与驾驶舱接口同源），"
        "排除 channel=admin_preview 的管理员自测会话；"
        "无可识别数据时比例返回 0、条数返回 0，不使用估算值。\n"
        "导出格式：Markdown（由前端按本接口返回的 sections 生成）。"
    )

    return [
        {"title": "一、本周摘要", "body": None, "metrics": summary_metrics},
        {"title": "二、北极星与统计口径", "body": north_body, "metrics": north_metrics},
        {"title": "三、使用情况", "body": usage_body, "metrics": usage_metrics},
        {"title": "四、回答质量", "body": quality_body, "metrics": quality_metrics},
        {"title": "五、知识治理", "body": governance_body, "metrics": governance_metrics},
        {"title": "六、身份接入质量", "body": identity_body, "metrics": identity_metrics},
        {"title": "七、下周建议", "body": advice_body, "metrics": []},
        {"title": "八、口径与限制", "body": limitation_body, "metrics": []},
    ]


def _weekly_markdown(
    *, week_label: str, range_label: str, sections: list[dict[str, Any]], exported_at: str
) -> str:
    """周报 Markdown：由服务端同一批聚合直接生成，数字与看板/JSON 不会漂移。

    注意：契约 `WeeklyReport` 的正式响应类型是 JSON（sections[].metrics 是数据源）；
    这里额外返回 `markdown` 字段供前端"复制/下载 .md"直接使用，属于附加字段，
    不替代契约要求的 JSON 结构。
    """
    lines = [
        "# 售后知识助手 · 周报",
        "",
        f"> 统计周 {week_label} · 区间 {range_label} · 生成时间 {exported_at}",
        "> 数据来源：服务端真实库聚合（与驾驶舱同源；排除 channel=admin_preview）",
    ]
    for section in sections:
        lines += ["", f"## {section['title']}", ""]
        for metric in section.get("metrics") or []:
            value = metric.get("value")
            unit = metric.get("unit") or ""
            if isinstance(value, int | float) and unit == "%":
                value_text = f"{value}%"
            elif isinstance(value, int | float) and unit:
                value_text = f"{value} {unit}"
            else:
                value_text = str(value)
            note_parts = []
            if isinstance(metric.get("delta"), int | float):
                delta = metric["delta"]
                if unit in ("%", "pt"):
                    note_parts.append(f"环比 {delta:+.1f} pt")
                else:
                    note_parts.append(f"环比 {delta:+g}{(' ' + unit) if unit else ''}")
            if metric.get("source"):
                note_parts.append(f"来源：{metric['source']}")
            suffix = f"（{'；'.join(note_parts)}）" if note_parts else ""
            lines.append(f"- **{metric['label']}**：{value_text}{suffix}")
        body = section.get("body")
        if body:
            lines.append("")
            for body_line in str(body).split("\n"):
                if not body_line.strip():
                    continue
                ordered = body_line.startswith(("1.", "2.", "3.", "4.", "5."))
                lines.append(body_line if ordered else f"- {body_line}")
    lines += [
        "",
        "---",
        "",
        "本报告由服务端聚合接口生成；未采集的指标留空，不使用估算值。",
    ]
    return "\n".join(lines) + "\n"


@router.get("/reports/weekly")
def weekly_report(
    week: str = Query(..., description="ISO 周，如 2026-W42"),
    conn: psycopg.Connection = Depends(get_conn),
    settings: Settings = Depends(get_settings),
    admin: dict = Depends(require_governance_admin),
) -> dict[str, Any]:
    """周报八节：所有数字复用看板同一批聚合（前端据此生成 Markdown 下载）。"""
    week_label, start, end = _parse_week(week)
    prev_start = start - timedelta(days=7)
    with conn.cursor(row_factory=dict_row) as cur:
        totals = _round_totals(cur, start, end)
        prev_totals = _round_totals(cur, prev_start, start)
        sessions = _session_counts(cur, start, end)
        kpi = _week_kpi(totals, _handoff_sessions(cur, start, end))
        prev_kpi = _week_kpi(prev_totals, _handoff_sessions(cur, prev_start, start))
        sections = _build_weekly_sections(
            totals=totals,
            kpi=kpi,
            prev_kpi=prev_kpi,
            sessions=sessions,
            cur=cur,
            start=start,
            end=end,
            settings=settings,
        )
    first_week = start - timedelta(days=35)
    first_iso = first_week.isocalendar()
    range_label = f"{first_iso.year}-W{first_iso.week:02d} ~ {week_label}"
    exported_at = datetime.now(UTC).isoformat(timespec="seconds")
    return {
        "week": week_label,
        "range": range_label,
        "sections": sections,
        "exportedAt": exported_at,
        "format": "markdown",
        # 附加字段（契约 WeeklyReport 之外的便利字段）：服务端生成的 Markdown 文本
        "markdown": _weekly_markdown(
            week_label=week_label,
            range_label=range_label,
            sections=sections,
            exported_at=exported_at,
        ),
    }
