"""Excel FAQ 解析：一行一条问答，直接生成证据单元（FAQ 是原子，不再按字数切）。

表格约定（表头大小写/空格不敏感，支持中英文别名）：
- 问题 / question / q（必填）
- 答案 / answer / a（必填）
- 分类 / category（可选，写入 heading 与证据元数据）
- 关键词 / keywords（可选，用｜、；、顿号分隔，写入元数据供检索判断）
- 行级批次（可选）：**列名由调用方注入**（`batch_keys`），解析器不认识行业词；
  有值的行写入 `ChunkDraft.batch`，作为该证据单元的批次覆盖；空值留给文档级批次。
- 其它列忽略，不猜测含义。

出处定位用「工作表名 + 行号」，不伪造成 PDF 页码；`page_from/page_to` 置 0，
由序列化层显示为“无页码”，loc 形如 `FAQ 第2行`。
"""

from __future__ import annotations

import re
from collections.abc import Iterable
from pathlib import Path
from typing import Any

from .types import ChunkDraft

_QUESTION_KEYS = {"问题", "question", "q", "标准问题", "用户问题"}
_ANSWER_KEYS = {"答案", "answer", "a", "回答", "标准答案"}
_CATEGORY_KEYS = {"分类", "类别", "category", "类型"}
_KEYWORD_KEYS = {"关键词", "关键字", "keywords", "keyword"}
_HEADER_SCAN_ROWS = 20
_MAX_ROWS = 2000
_MAX_TEXT_CHARS = 6000


def _key(raw: Any) -> str:
    return str(raw or "").strip().lower().replace(" ", "").replace("\u3000", "")


def _cell(values: Any, index: int | None) -> str:
    if index is None or index >= len(values):
        return ""
    value = values[index]
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _split_keywords(raw: str) -> list[str]:
    words: list[str] = []
    for part in re.split(r"[｜|;；、,，\n]", raw or ""):
        word = part.strip()
        if word and word not in words:
            words.append(word)
    return words


def _find_header(
    rows: list[list[Any]], batch_keys: frozenset[str]
) -> tuple[int, dict[str, int]]:
    for index, row in enumerate(rows[:_HEADER_SCAN_ROWS]):
        mapping: dict[str, int] = {}
        for column, value in enumerate(row):
            key = _key(value)
            if key in _QUESTION_KEYS:
                mapping.setdefault("question", column)
            elif key in _ANSWER_KEYS:
                mapping.setdefault("answer", column)
            elif key in _CATEGORY_KEYS:
                mapping.setdefault("category", column)
            elif key in _KEYWORD_KEYS:
                mapping.setdefault("keywords", column)
            elif key in batch_keys:
                mapping.setdefault("batch", column)
        if "question" in mapping and "answer" in mapping:
            return index, mapping
    raise ValueError("FAQ Excel 缺少表头：至少需要「问题」「答案」两列")


def parse_faq_xlsx(
    path: str | Path,
    *,
    batch_keys: Iterable[str] = (),
) -> tuple[list[ChunkDraft], dict[str, Any]]:
    """解析 FAQ 工作表；任何一行缺问题/答案都整批失败，不静默跳过。

    `batch_keys`：可选「批次」列的列名别名（行业词汇由 app 层传入）。
    传空表示不识别行级批次（默认行为，与旧版一致）。
    """
    import openpyxl  # 延迟导入：只有解析 FAQ 时才需要

    normalized_batch_keys = frozenset(_key(alias) for alias in batch_keys if _key(alias))

    workbook = openpyxl.load_workbook(filename=path, read_only=True, data_only=True)
    try:
        sheet_name = "FAQ" if "FAQ" in workbook.sheetnames else workbook.sheetnames[0]
        sheet = workbook[sheet_name]
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()

    header_index, columns = _find_header(rows, normalized_batch_keys)
    drafts: list[ChunkDraft] = []
    errors: list[str] = []
    for row_no, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        question = _cell(row, columns.get("question"))
        answer = _cell(row, columns.get("answer"))
        if not question and not answer:
            continue
        if not question or not answer:
            errors.append(f"第 {row_no} 行「问题」和「答案」都必须填写")
            continue
        category = _cell(row, columns.get("category"))
        keywords = _split_keywords(_cell(row, columns.get("keywords")))
        # 行级批次：有值才覆盖文档批次；空值返回 None，交给文档级批次兜底
        row_batch = _cell(row, columns.get("batch")) or None
        # 正文以「问题」开头：切片列表只显示前两行，问题才是这条证据的身份；
        # 分类已经在 heading 与 meta 里，正文末尾再留一份便于阅读与检索。
        text = f"问题：{question}\n答案：{answer}"
        if category:
            text = f"{text}\n分类：{category}"
        if len(text) > _MAX_TEXT_CHARS:
            errors.append(f"第 {row_no} 行问答超过 {_MAX_TEXT_CHARS} 字，请拆成多条 FAQ")
            continue
        drafts.append(
            ChunkDraft(
                seq=len(drafts) + 1,
                text=text,
                loc=f"{sheet_name} 第{row_no}行",
                page_from=0,
                page_to=0,
                heading=category or "FAQ",
                batch=row_batch,
                meta={
                    "sheet": sheet_name,
                    "row": row_no,
                    "category": category or None,
                    "keywords": keywords,
                },
            )
        )
    if errors:
        head = "；".join(errors[:5])
        more = f"；另有 {len(errors) - 5} 行" if len(errors) > 5 else ""
        raise ValueError(head + more)
    if not drafts:
        raise ValueError("FAQ Excel 里没有可入库的问答")
    if len(drafts) > _MAX_ROWS:
        raise ValueError(f"FAQ 单次最多 {_MAX_ROWS} 条，请拆分文件")
    return drafts, {"rows": len(drafts), "sheet": sheet_name}
